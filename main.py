import asyncio
import json
import os
import re
import shutil
import unicodedata
import uuid
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import List
from zoneinfo import ZoneInfo

from dotenv import load_dotenv
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from docx import Document
from num2words import num2words

from models import CAMPOS_ADMISSAO, DadosAdmissao, GerarArquivosResponse

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent
GENERATED_DIR = BASE_DIR / "generated"
GENERATED_DIR.mkdir(exist_ok=True)
TEMPLATE_PATH = BASE_DIR / "template_contrato.docx"

BRASILIA_TZ = ZoneInfo("America/Sao_Paulo")

# Estado em memória dos jobs de geração (sem banco de dados, conforme escopo do projeto).
# Guarda apenas o nome-base usado nos arquivos baixáveis. Os arquivos em si ficam no disco
# até a limpeza diária (ver _limpeza_diaria_loop) apagar tudo de uma vez, de madrugada.
JOBS_META: dict[str, dict] = {}


def _slugify_nome(nome: str) -> str:
    sem_acentos = unicodedata.normalize("NFKD", nome).encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^a-zA-Z0-9]+", "", sem_acentos).lower()
    return slug or "colaborador"


def _limpar_todos_os_jobs() -> None:
    for job_dir in GENERATED_DIR.iterdir():
        if job_dir.is_dir():
            shutil.rmtree(job_dir, ignore_errors=True)
    JOBS_META.clear()


async def _limpeza_diaria_loop() -> None:
    """Apaga todos os arquivos gerados uma vez por dia, às 3h no horário de Brasília.

    Não apagamos na hora do download: os arquivos ficam disponíveis no servidor
    durante o dia (para o RH baixar/reenviar quando precisar) e são varridos de
    uma vez só de madrugada, quando ninguém deve estar usando o sistema.
    """
    while True:
        agora = datetime.now(BRASILIA_TZ)
        proxima_execucao = agora.replace(hour=3, minute=0, second=0, microsecond=0)
        if proxima_execucao <= agora:
            proxima_execucao += timedelta(days=1)
        await asyncio.sleep((proxima_execucao - agora).total_seconds())
        _limpar_todos_os_jobs()


ANTHROPIC_MODEL = os.environ.get("ANTHROPIC_MODEL", "claude-sonnet-5")

app = FastAPI(title="Admissão RH — API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
async def _iniciar_limpeza_diaria() -> None:
    asyncio.create_task(_limpeza_diaria_loop())

# --------------------------------------------------------------------------
# Extração de dados via IA (Anthropic)
# --------------------------------------------------------------------------

SYSTEM_PROMPT = """Você é um assistente especialista em extrair dados cadastrais de documentos brasileiros a partir de imagens ou PDFs, para fins de admissão de funcionários.

Você pode receber mais de um documento na mesma solicitação — por exemplo, um documento de identificação (RG, CPF ou CNH) e, separadamente, um comprovante de residência (conta de luz, água, telefone ou internet). Combine as informações de TODOS os documentos enviados num único resultado: dados pessoais normalmente vêm do documento de identificação, e o endereço normalmente vem do comprovante de residência (mas também pode vir do próprio documento de identificação, quando presente).

Analise cuidadosamente os documentos fornecidos e devolva a resposta APENAS como um objeto JSON válido — sem markdown, sem texto antes ou depois, sem comentários — contendo exatamente estas chaves:

{
  "nome_completo": "",
  "cpf": "",
  "rg": "",
  "rg_orgao_emissor": "",
  "rg_uf_emissor": "",
  "rg_data_emissao": "",
  "data_nascimento": "",
  "sexo": "",
  "naturalidade_cidade": "",
  "naturalidade_uf": "",
  "nome_mae": "",
  "nome_pai": "",
  "endereco_logradouro": "",
  "endereco_numero": "",
  "endereco_complemento": "",
  "endereco_bairro": "",
  "endereco_cidade": "",
  "endereco_uf": "",
  "endereco_cep": ""
}

Regras:
- Nunca invente ou deduza informação que não esteja visível nos documentos. Se um campo não for encontrado em nenhum documento, use string vazia "".
- Formate o CPF como 000.000.000-00 quando encontrado.
- Formate datas como DD/MM/AAAA.
- "sexo" deve ser "Masculino" ou "Feminino" quando identificável.
- "rg_orgao_emissor" é a sigla do órgão emissor (ex: "SSP"), separada da UF, que vai em "rg_uf_emissor" (ex: "RJ").
- "naturalidade_cidade" e "naturalidade_uf" vêm separados (ex: cidade "Rio de Janeiro", uf "RJ").
- Os campos "endereco_*" vêm cada um separado (não junte tudo num só campo). "endereco_cep" apenas números.
- Responda SOMENTE com o objeto JSON."""

IMAGE_MEDIA_TYPES = {"image/png", "image/jpeg", "image/jpg", "image/webp", "image/gif"}


def _build_content_block(content_type: str, data_b64: str) -> dict:
    if content_type == "application/pdf":
        return {
            "type": "document",
            "source": {"type": "base64", "media_type": "application/pdf", "data": data_b64},
        }
    media_type = content_type if content_type in IMAGE_MEDIA_TYPES else "image/jpeg"
    return {
        "type": "image",
        "source": {"type": "base64", "media_type": media_type, "data": data_b64},
    }


def _extract_json(text: str) -> dict:
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass
    match = re.search(r"\{.*\}", text, re.DOTALL)
    if match:
        return json.loads(match.group(0))
    raise ValueError("A IA não retornou um JSON válido.")


@app.post("/api/extrair")
async def extrair_dados(documentos: List[UploadFile] = File(...)):
    if not documentos:
        raise HTTPException(status_code=400, detail="Nenhum documento enviado.")

    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        raise HTTPException(
            status_code=503,
            detail=(
                "ANTHROPIC_API_KEY não configurada no servidor. Defina a variável "
                "de ambiente (ou arquivo .env) para habilitar a extração via IA."
            ),
        )

    import base64

    import anthropic

    content_blocks = []
    for upload in documentos:
        raw = await upload.read()
        if not raw:
            continue
        b64 = base64.b64encode(raw).decode("utf-8")
        content_blocks.append(_build_content_block(upload.content_type or "", b64))

    if not content_blocks:
        raise HTTPException(status_code=400, detail="Os arquivos enviados estão vazios.")

    content_blocks.append(
        {
            "type": "text",
            "text": "Extraia os dados cadastrais dos documentos de identificação acima.",
        }
    )

    client = anthropic.Anthropic(api_key=api_key)

    try:
        response = client.messages.create(
            model=ANTHROPIC_MODEL,
            max_tokens=1024,
            system=SYSTEM_PROMPT,
            messages=[{"role": "user", "content": content_blocks}],
        )
    except anthropic.APIError as exc:
        raise HTTPException(status_code=502, detail=f"Erro ao chamar a API da Anthropic: {exc}")

    texto_resposta = "".join(
        block.text for block in response.content if getattr(block, "type", None) == "text"
    )

    try:
        dados_brutos = _extract_json(texto_resposta)
    except (ValueError, json.JSONDecodeError):
        raise HTTPException(
            status_code=502,
            detail="A IA retornou uma resposta que não pôde ser interpretada como JSON.",
        )

    dados = DadosAdmissao(**{k: dados_brutos.get(k, "") or "" for k in CAMPOS_ADMISSAO})
    return dados.model_dump()


# --------------------------------------------------------------------------
# Geração de arquivos (contrato .docx + layout de importação do Domínio)
# --------------------------------------------------------------------------


def _replace_placeholders_in_paragraph(paragraph, contexto: dict) -> None:
    if not paragraph.runs:
        return
    texto_completo = "".join(run.text for run in paragraph.runs)
    if "{{" not in texto_completo:
        return
    novo_texto = texto_completo
    for chave, valor in contexto.items():
        novo_texto = novo_texto.replace("{{" + chave + "}}", str(valor))
    if novo_texto == texto_completo:
        return
    paragraph.runs[0].text = novo_texto
    for run in paragraph.runs[1:]:
        run.text = ""


def _juntar(*partes: str, sep: str = " ") -> str:
    return sep.join(p for p in partes if p)


ESTADO_CIVIL_EXTENSO = {
    "S": "solteiro(a)",
    "C": "casado(a)",
    "V": "viúvo(a)",
    "D": "divorciado(a)",
    "O": "em regime de concubinato",
    "J": "separado(a) judicialmente",
    "U": "em união estável",
}

MESES_PT = [
    "", "janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho",
    "agosto", "setembro", "outubro", "novembro", "dezembro",
]


def _formatar_moeda_brl(valor_str: str) -> str:
    try:
        valor = Decimal(str(valor_str).replace(",", "."))
    except (InvalidOperation, ValueError):
        return ""
    texto = f"{valor:,.2f}"
    return texto.replace(",", "_").replace(".", ",").replace("_", ".")


def _valor_por_extenso_reais(valor_str: str) -> str:
    try:
        valor = Decimal(str(valor_str).replace(",", "."))
    except (InvalidOperation, ValueError):
        return ""
    reais = int(valor)
    centavos = int(round((valor - reais) * 100))
    partes = [f"{num2words(reais, lang='pt_BR')} {'real' if reais == 1 else 'reais'}"]
    if centavos:
        partes.append(f"{num2words(centavos, lang='pt_BR')} {'centavo' if centavos == 1 else 'centavos'}")
    return " e ".join(partes)


def _montar_campos_combinados(dados: DadosAdmissao) -> dict:
    """Strings de exibição (para o contrato), combinando os campos estruturados."""
    endereco_linha1 = _juntar(
        dados.endereco_logradouro,
        f", {dados.endereco_numero}" if dados.endereco_numero else "",
        sep="",
    )
    endereco_completo = " - ".join(
        p
        for p in [
            endereco_linha1,
            dados.endereco_complemento,
            dados.endereco_bairro,
            _juntar(dados.endereco_cidade, dados.endereco_uf, sep="/"),
            f"CEP {dados.endereco_cep}" if dados.endereco_cep else "",
        ]
        if p
    )
    hoje = datetime.now(BRASILIA_TZ)
    return {
        "endereco_completo": endereco_completo,
        "naturalidade": _juntar(dados.naturalidade_cidade, dados.naturalidade_uf, sep="/"),
        "rg_orgao_uf": _juntar(dados.rg_orgao_emissor, dados.rg_uf_emissor, sep="/"),
        "estado_civil_extenso": ESTADO_CIVIL_EXTENSO.get(dados.estado_civil, ""),
        "salario_formatado": _formatar_moeda_brl(dados.salario),
        "salario_extenso": _valor_por_extenso_reais(dados.salario),
        "dia_assinatura": str(hoje.day),
        "mes_assinatura": MESES_PT[hoje.month],
        "ano_assinatura": str(hoje.year),
    }


def gerar_contrato_docx(dados: DadosAdmissao, destino: Path) -> None:
    if not TEMPLATE_PATH.exists():
        raise HTTPException(
            status_code=500,
            detail=(
                f"Template não encontrado em {TEMPLATE_PATH.name}. Execute "
                "scripts/create_template.py ou adicione o modelo real da empresa."
            ),
        )
    doc = Document(str(TEMPLATE_PATH))
    contexto = {**dados.model_dump(), **_montar_campos_combinados(dados)}

    for paragraph in doc.paragraphs:
        _replace_placeholders_in_paragraph(paragraph, contexto)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_placeholders_in_paragraph(paragraph, contexto)

    doc.save(str(destino))


with open(BASE_DIR / "data" / "municipios_dominio.json", encoding="utf-8") as _f:
    _MUNICIPIOS_DOMINIO: dict[str, int] = json.load(_f)

PAIS_BRASIL_CODIGO = "30"  # Código do Brasil na aba "Código dos países" do Domínio.


def _codigo_municipio(nome_cidade: str) -> str:
    """Busca exata (sem acento/caixa) na tabela oficial de municípios do Domínio.

    Só preenche quando encontra correspondência exata — na dúvida, deixa em
    branco para o RH completar manualmente, em vez de arriscar um código errado.
    """
    if not nome_cidade:
        return ""
    chave = unicodedata.normalize("NFKD", nome_cidade).encode("ascii", "ignore").decode("ascii").strip().upper()
    codigo = _MUNICIPIOS_DOMINIO.get(chave)
    return str(codigo) if codigo is not None else ""


def _somente_digitos(valor: str) -> str:
    return re.sub(r"\D", "", valor or "")


def _mapear_sexo_dominio(sexo: str) -> str:
    s = (sexo or "").strip().lower()
    if s.startswith("m"):
        return "M"
    if s.startswith("f"):
        return "F"
    return ""


# Posição (1-indexado, conforme o "Dicionário dos campos" do Domínio) -> valor.
# Campos documentais vêm da extração por IA; campos de folha de pagamento vêm
# do preenchimento manual do RH na tela de validação (ambos no mesmo objeto
# DadosAdmissao — ver models.py).
def _montar_linha_dominio(dados: DadosAdmissao) -> list[str]:
    campos: dict[int, str] = {
        1: "0",  # Código empresa — instrução do próprio Domínio: sempre 0.
        2: dados.codigo_empregado,
        3: dados.nome_completo,
        4: _somente_digitos(dados.cpf),
        5: dados.codigo_filial,
        6: dados.codigo_servico,
        7: dados.codigo_cargo,
        8: dados.codigo_departamento,
        9: dados.codigo_centro_custos,
        10: dados.salario,
        11: dados.data_admissao,
        12: dados.codigo_esocial_empregado,
        13: dados.vinculo_empregaticio,
        14: dados.optante_fgts,
        15: dados.data_opcao_fgts,
        16: dados.ocorrencia_sefip,
        17: dados.endereco_logradouro,
        18: dados.endereco_numero,
        19: dados.endereco_complemento,
        20: dados.endereco_bairro,
        21: dados.endereco_cidade,
        22: dados.endereco_uf,
        23: _somente_digitos(dados.endereco_cep),
        24: dados.data_nascimento,
        25: dados.codigo_banco,
        26: dados.conta_corrente,
        27: dados.digito_conta,
        28: dados.tipo_conta,
        29: dados.codigo_convenio_coletivo,
        30: dados.multiplos_vinculos,
        31: dados.pagou_contribuicao_sindical,
        32: dados.vencimento_ferias,
        33: dados.horas_dia,
        34: dados.categoria,
        35: dados.tipo_indicativo_admissao,
        36: dados.grau_instrucao,
        37: dados.deficiente_fisico,
        38: dados.tipo_admissao,
        39: dados.horas_semana,
        40: dados.forma_pagamento,
        41: dados.horas_mes,
        42: dados.situacao_qualificacao_cadastral,
        43: _mapear_sexo_dominio(dados.sexo),
        44: PAIS_BRASIL_CODIGO,  # País nacionalidade — assumimos Brasil por padrão.
        45: dados.raca_cor,
        46: dados.estado_civil,
        47: dados.recebia_seguro_desemprego,
        48: dados.codigo_jornada,
        49: dados.nome_pai,
        50: dados.nome_mae,
        51: dados.tipo_endereco,
        52: _codigo_municipio(dados.endereco_cidade),
        53: _codigo_municipio(dados.naturalidade_cidade),
        54: PAIS_BRASIL_CODIGO,  # País do endereço.
        55: PAIS_BRASIL_CODIGO,  # País de nascimento.
        # Na CIN (nova identidade unificada) o número do RG às vezes não existe —
        # nesse caso o CPF passa a ser o próprio número de identidade. Só cai para o
        # CPF quando o RG realmente não veio de nenhum documento.
        56: dados.rg.strip() or dados.cpf,
        57: dados.rg_orgao_emissor,
        58: dados.rg_uf_emissor,
        59: dados.rg_data_emissao,
        60: dados.contrato_experiencia,
        61: dados.clausula_assecuratoria,
        62: dados.dias_contrato_experiencia,
        63: dados.inicio_prazo_determinado,
        64: dados.fim_prazo_determinado,
        65: dados.dias_prorrogacao,
        66: dados.prorrogacao_prazo_determinado,
        67: dados.motivo_prorrogacao,
        68: dados.forma_calculo_pagamento,
        69: dados.atividade_simples_nacional,
        70: dados.email,
        71: dados.ddd_telefone,
        72: dados.telefone,
        73: dados.ddd_contato_2,
        74: dados.contato_2,
        75: dados.data_admissao,  # Data Vantagem — o layout exige a mesma data da admissão.
        76: dados.naturalidade_uf,
        77: dados.codigo_categoria_esocial,
        78: dados.codigo_sindicato_cadastro,
    }
    return [campos.get(pos, "") for pos in range(1, 79)]


def gerar_layout_dominio(dados: DadosAdmissao, destino: Path) -> None:
    """Gera o "empregados.txt" no layout oficial de importação de admissão do
    Domínio Web: 78 colunas separadas por TAB, sem cabeçalho, uma linha por
    empregado (mesmo formato que a macro da planilha modelo do escritório gera).

    Campos documentais vêm da extração por IA; campos de folha de pagamento
    vêm do preenchimento manual do RH. O que não for preenchido em nenhum dos
    dois fica em branco, para o RH completar dentro do próprio Domínio Web.
    """
    linha = "\t".join(_montar_linha_dominio(dados))
    with destino.open("w", encoding="cp1252", errors="replace", newline="\r\n") as f:
        f.write(linha + "\n")


# Rótulos curtos das 78 posições (mesma numeração do "Dicionário dos campos"),
# usados só no arquivo de conferência legível — o .txt oficial não tem rótulos.
FIELD_LABELS_DOMINIO = [
    "Código empresa", "Código empregado", "Nome", "CPF", "Código filial",
    "Código serviço", "Código cargo", "Código departamento", "Centro de custos",
    "Salário", "Data admissão", "Código eSocial", "Vínculo empregatício",
    "Optante FGTS", "Data opção FGTS", "Ocorrência SEFIP", "Endereço",
    "Número endereço", "Complemento", "Bairro", "Cidade", "UF", "CEP",
    "Data nascimento", "Código banco", "Conta corrente", "Dígito conta",
    "Tipo de conta", "Convenção coletiva", "Múltiplos vínculos",
    "Pagou contr. sindical", "Vencimento férias", "Horas dia", "Categoria",
    "Tipo indicativo admissão", "Grau instrução", "Deficiente físico",
    "Tipo de admissão", "Horas semana", "Forma pagamento", "Horas mês",
    "Situação qualificação cadastral", "Sexo", "País nacionalidade",
    "Raça/Cor", "Estado civil", "Recebia seguro desemprego", "Código jornada",
    "Nome do pai", "Nome da mãe", "Tipo de endereço",
    "Código município endereço", "Código município nascimento",
    "Código país endereço", "Código país nascimento", "Identidade (RG)",
    "Órgão expedição RG", "UF expedição RG", "Data expedição RG",
    "Contrato de experiência", "Cláusula assecuratória",
    "Dias contrato experiência", "Início prazo determinado",
    "Fim prazo determinado", "Dias prorrogação",
    "Prorrogação prazo determinado", "Motivo prorrogação",
    "Forma pagamento (mês/semana)", "Atividade simples nacional", "E-mail",
    "DDD telefone", "Telefone", "DDD contato 2", "Contato 2", "Data Vantagem",
    "UF nascimento", "Código categoria eSocial", "Código sindicato",
]


def gerar_conferencia_dominio(dados: DadosAdmissao, destino: Path) -> None:
    """Versão legível (com rótulos) das mesmas 78 posições do arquivo oficial,
    só para o RH conferir visualmente antes de importar — não é o arquivo que
    vai pro Domínio (esse é o layout_dominio.txt, cru e sem rótulos).
    """
    valores = _montar_linha_dominio(dados)
    linhas = [
        f"{pos:02d} - {label}: {valor}"
        for pos, (label, valor) in enumerate(zip(FIELD_LABELS_DOMINIO, valores), start=1)
        if valor
    ]
    destino.write_text("\n".join(linhas) + "\n", encoding="utf-8")


@app.post("/api/gerar", response_model=GerarArquivosResponse)
async def gerar_arquivos(dados: DadosAdmissao):
    if not dados.nome_completo.strip():
        raise HTTPException(status_code=400, detail="nome_completo é obrigatório.")

    job_id = uuid.uuid4().hex
    job_dir = GENERATED_DIR / job_id
    job_dir.mkdir(parents=True, exist_ok=True)

    docx_path = job_dir / "contrato_trabalho.docx"
    dominio_path = job_dir / "layout_dominio.txt"
    conferencia_path = job_dir / "conferencia_dominio.txt"

    gerar_contrato_docx(dados, docx_path)
    gerar_layout_dominio(dados, dominio_path)
    gerar_conferencia_dominio(dados, conferencia_path)

    # Nome-base usado nos arquivos baixados: {nomedocolaborador}_{ddmmyyyy} (data de
    # geração dos documentos — a mesma convenção que será usada na pasta do cloud).
    base_nome = f"{_slugify_nome(dados.nome_completo)}_{date.today().strftime('%d%m%Y')}"

    JOBS_META[job_id] = {"base_nome": base_nome}

    return GerarArquivosResponse(
        job_id=job_id,
        docx_url=f"/api/download/{job_id}/docx",
        dominio_url=f"/api/download/{job_id}/dominio",
        conferencia_url=f"/api/download/{job_id}/conferencia",
    )


@app.get("/api/download/{job_id}/{tipo}")
async def download_arquivo(job_id: str, tipo: str):
    if not re.fullmatch(r"[a-f0-9]{32}", job_id):
        raise HTTPException(status_code=400, detail="job_id inválido.")

    meta = JOBS_META.get(job_id)
    if not meta:
        raise HTTPException(status_code=404, detail="Arquivo não encontrado (job expirado ou já apagado na limpeza diária).")

    job_dir = GENERATED_DIR / job_id
    if tipo == "docx":
        caminho = job_dir / "contrato_trabalho.docx"
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        nome_arquivo = f"{meta['base_nome']}_contrato.docx"
    elif tipo == "dominio":
        caminho = job_dir / "layout_dominio.txt"
        media_type = "text/plain"
        nome_arquivo = f"{meta['base_nome']}_dominio.txt"
    elif tipo == "conferencia":
        caminho = job_dir / "conferencia_dominio.txt"
        media_type = "text/plain"
        nome_arquivo = f"{meta['base_nome']}_conferencia.txt"
    else:
        raise HTTPException(status_code=404, detail="Tipo de arquivo desconhecido.")

    if not caminho.exists():
        raise HTTPException(status_code=404, detail="Arquivo não encontrado (job expirado ou inválido).")

    return FileResponse(caminho, media_type=media_type, filename=nome_arquivo)


# --------------------------------------------------------------------------
# Frontend estático
# --------------------------------------------------------------------------

app.mount("/", StaticFiles(directory=str(BASE_DIR / "static"), html=True), name="static")
