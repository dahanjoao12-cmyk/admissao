"""Gera um template_contrato.docx MOCK com placeholders {{campo}} para testes.

Em produção, o RH deve substituir este arquivo pelo modelo real de contrato da
empresa (mesmos placeholders, ou ajustar CAMPOS_ADMISSAO / main.py conforme o
modelo real).
"""
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

sys.path.append(str(Path(__file__).resolve().parent.parent))


def build_template(output_path: Path) -> None:
    doc = Document()

    title = doc.add_heading("CONTRATO INDIVIDUAL DE TRABALHO", level=1)
    title.alignment = 1

    doc.add_paragraph(
        "Pelo presente instrumento particular, de um lado a empresa "
        "[RAZÃO SOCIAL DA EMPRESA], e de outro lado o(a) empregado(a) "
        "abaixo qualificado(a), têm entre si justo e contratado o "
        "seguinte:"
    )

    doc.add_heading("1. Qualificação do Empregado", level=2)

    campos = [
        ("Nome completo:", "{{nome_completo}}"),
        ("CPF:", "{{cpf}}"),
        ("RG:", "{{rg}}"),
        ("Órgão emissor:", "{{rg_orgao_uf}}"),
        ("Data de nascimento:", "{{data_nascimento}}"),
        ("Sexo:", "{{sexo}}"),
        ("Naturalidade:", "{{naturalidade}}"),
        ("Nome da mãe:", "{{nome_mae}}"),
        ("Nome do pai:", "{{nome_pai}}"),
        ("Endereço:", "{{endereco_completo}}"),
    ]

    for label, placeholder in campos:
        p = doc.add_paragraph()
        run_label = p.add_run(f"{label} ")
        run_label.bold = True
        run_label.font.size = Pt(11)
        run_value = p.add_run(placeholder)
        run_value.font.size = Pt(11)

    doc.add_heading("2. Cláusulas", level=2)
    doc.add_paragraph(
        "O(A) empregado(a) acima qualificado(a) fica admitido(a) nos "
        "termos da legislação trabalhista vigente, comprometendo-se "
        "ambas as partes a cumprir fielmente as condições aqui "
        "estabelecidas."
    )

    doc.add_paragraph("\n\n_________________________________")
    doc.add_paragraph("Assinatura do Empregador")
    doc.add_paragraph("\n\n_________________________________")
    doc.add_paragraph("Assinatura do(a) Empregado(a): {{nome_completo}}")

    doc.save(str(output_path))
    print(f"Template criado em: {output_path}")


if __name__ == "__main__":
    root = Path(__file__).resolve().parent.parent
    build_template(root / "template_contrato.docx")
