"""Prepara o contrato real do escritório para uso pelo sistema: troca os
placeholders em colchetes ([Nome do Empregado], [Inserir CPF] etc.) pelos
tokens {{campo}} que o gerar_contrato_docx() do main.py substitui.

Roda uma única vez sobre o arquivo original (nunca sobrescreve o original —
lê de ORIGEM e grava em template_contrato.docx, na raiz do projeto).

Uso: python scripts/preparar_template_real.py "<caminho do docx original>"
"""
import sys
from pathlib import Path

from docx import Document

# Paragraph index -> novo texto (com {{tokens}}). Descoberto inspecionando o
# arquivo original enviado pelo escritório (02/09/2026). Se o escritório
# editar o contrato de novo, os índices abaixo podem não bater mais — rode
# scripts/create_template.py --dump (ou releia com python-docx) para
# conferir antes de reaplicar.
SUBSTITUICOES = {
    8: (
        "EMPREGADO(A): {{nome_completo}}, {{nacionalidade}}, {{estado_civil_extenso}}, "
        "portador(a) da Cédula de Identidade RG nº {{rg}}, inscrito(a) no CPF sob o nº "
        "{{cpf}}, titular da CTPS nº {{numero_ctps}} série {{serie_ctps}}, residente e "
        "domiciliado(a) na {{endereco_logradouro}}, nº {{endereco_numero}}, Cidade "
        "{{endereco_cidade}} – {{endereco_uf}}."
    ),
    17: (
        " O(A) EMPREGADO(A) prestará seus serviços junto ao EMPREGADOR no exercício da "
        "função de {{funcao}}."
    ),
    25: (
        "O(A) EMPREGADO(A) receberá uma remuneração mensal equivalente a R$ "
        "{{salario_formatado}} ({{salario_extenso}}), que será pago até o 5º dia útil de "
        "cada mês."
    ),
    29: (
        "O EMPREGADOR fornecerá ao(a) EMPREGADO(A), ajuda de custo para alimentação, "
        "mediante desconto mensal de R$ {{desconto_alimentacao}}, que pode ser "
        "anualmente ajustado de acordo com os índices da inflação, devidamente "
        "autorizado neste ato pelo(a) EMPREGADO(A), possuindo caráter eminentemente "
        "indenizatório e não salarial, nos termos do art. 457, § 2º da CLT."
    ),
    37: (
        "O(A) EMPREGADO(A) exercerá suas atividades presenciais, na sede do EMPREGADOR "
        "localizado Avenida Rio Branco, nº 99, 5º andar, Centro, Rio de Janeiro – RJ, "
        "com expediente diário no horário compreendido de {{horario_seg_qui}}, de "
        "Segunda-feira à Quinta-feira, e no horário de {{horario_sexta}} às "
        "sextas-feiras, como acordo de compensação mensal de jornada entre EMPREGADO(A) "
        "E EMPREGADOR, cumprindo e totalizando 44 horas semanais conforme previsão na "
        "CLT."
    ),
    67: "Rio de Janeiro, {{dia_assinatura}} de {{mes_assinatura}} de {{ano_assinatura}}.",
    71: "{{representante_legal}}\nEmpregador",
    74: "{{nome_completo}}\nEmpregado(a)",
}


def preparar(origem: Path, destino: Path) -> None:
    doc = Document(str(origem))
    for indice, novo_texto in SUBSTITUICOES.items():
        paragrafo = doc.paragraphs[indice]
        texto_atual = paragrafo.text
        if not paragrafo.runs:
            raise RuntimeError(f"Parágrafo {indice} sem runs — não deveria acontecer.")
        paragrafo.runs[0].text = novo_texto
        for run in paragrafo.runs[1:]:
            run.text = ""
        print(f"[{indice}] {texto_atual!r}\n     -> {novo_texto!r}\n")
    doc.save(str(destino))
    print(f"Template salvo em: {destino}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Uso: python scripts/preparar_template_real.py \"<caminho do docx original>\"")
        raise SystemExit(1)
    origem = Path(sys.argv[1])
    destino = Path(__file__).resolve().parent.parent / "template_contrato.docx"
    preparar(origem, destino)
